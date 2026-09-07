import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(BASE_DIR, "docs")
SCREENSHOTS_DIR = os.path.join(DOCS_DIR, "screenshots")
DOCX_PATH = os.path.join(DOCS_DIR, "EcoSort_Campus_Prototype_Documentation.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_documentation_docx():
    doc = Document()

    # Page setup (margins 0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles
    DARK_GREEN = RGBColor(20, 83, 45)
    PRIMARY_GREEN = RGBColor(22, 163, 74)

    # Header / Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("EcoSort Campus: Smart Waste Segregation\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = DARK_GREEN

    sub = title.add_run("Working Prototype Technical Documentation\n")
    sub.font.name = "Arial"
    sub.font.size = Pt(14)
    sub.font.color.rgb = PRIMARY_GREEN

    tag = title.add_run('"Know it. Sort it. Sustain it."\n')
    tag.font.name = "Arial"
    tag.font.size = Pt(12)
    tag.font.italic = True
    tag.font.color.rgb = RGBColor(4, 120, 87)

    prog = title.add_run("AI for Sustainability Virtual Internship 2026 (AICTE / 1M1B / IBM SkillsBuild)")
    prog.font.name = "Arial"
    prog.font.size = Pt(10)
    prog.font.bold = True
    prog.font.color.rgb = RGBColor(71, 85, 105)

    hero_img = os.path.join(SCREENSHOTS_DIR, "01_homepage.png")
    if os.path.exists(hero_img):
        doc.add_picture(hero_img, width=Inches(6.0))
        p_cap = doc.add_paragraph("Figure: EcoSort Campus Main Application Interface")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True

    doc.add_page_break()

    # Section 1: Problem Definition
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("1. Problem Statement & Campus Waste Crisis")
    r1.font.color.rgb = DARK_GREEN

    p = doc.add_paragraph(
        "University campuses operate as micro-cities where thousands of students, faculty, and mess staff produce "
        "substantial volumes of mixed solid waste daily. In high-traffic locations such as cafeterias, lecture blocks, "
        "and hostels, waste segregation frequently fails.\n\n"
        "Project Baseline Problem Statement (from Project PPT):\n"
        "\"Over 60% of recyclable materials on campus end up in landfills due to cross-contamination in high-traffic bins.\"\n"
        "*(Note: Baseline problem observation from initial campus audit and proposal concept, not a measured municipal statistic).*"
    )

    doc.add_heading("UN Sustainable Development Goals (SDGs)", level=2)
    p_sdg = doc.add_paragraph()
    p_sdg.add_run("• Primary SDG 12: Responsible Consumption and Production (Target 12.5) — ").bold = True
    p_sdg.add_run("Substantially reduce waste generation through recycling, preventing clean recyclables from being ruined by organic contaminants.\n")
    p_sdg.add_run("• Secondary SDG 11: Sustainable Cities and Communities — ").bold = True
    p_sdg.add_run("Reduces institutional environmental footprints and relieves local municipal landfill pressure.")

    # Section 2: AI Solution
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("2. Proposed AI Solution & Dual-Engine Architecture")
    r2.font.color.rgb = DARK_GREEN

    doc.add_paragraph(
        "EcoSort Campus uses natural language processing (NLP) and zero-shot entity extraction powered by IBM Granite "
        "foundation models (watsonx.ai). The solution is implemented cleanly as a full-stack Next.js web application "
        "with an App Router API backend (/api/classify).\n\n"
        "Dual-Engine Reliability Strategy:\n"
        "1. Live IBM Granite API: Interacts with ibm/granite-3-3-8b-instruct via watsonx.ai REST endpoints when credentials are provided in .env.local.\n"
        "2. Rule-Based Fallback Engine: Provides deterministic classification across all four campus categories for seamless evaluation, offline demo resilience, and automated testing without exposing broken states."
    )

    # Section 3: Campus Bin System
    h3 = doc.add_heading(level=1)
    r3 = h3.add_run("3. Four-Bin Campus Classification System")
    r3.font.color.rgb = DARK_GREEN

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_titles = ["Bin Colour", "Category", "Target Materials", "Disposal Action"]
    for i, title_text in enumerate(hdr_titles):
        hdr_cells[i].text = title_text
        set_cell_background(hdr_cells[i], "14532D")
        for r in hdr_cells[i].paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8.5)

    bins_data = [
        ("GREEN (#16a34a)", "Wet / Compost", "Food scraps, oily samosa wrappers, chai cups, greasy parcels", "Tear clean dry margins; compost food/oil residue"),
        ("BLUE (#2563eb)", "Dry / Recyclable", "Clean PET plastic #1, aluminium cans, dry cardboard boxes", "Empty liquids completely; squash flat to save volume"),
        ("RED (#dc2626)", "Hazardous E-Waste", "9V/AA batteries, electronic scrap, lab chemicals, PCB parts", "Tape terminals with cello-tape; drop in e-waste box"),
        ("BLACK (#1c1917)", "Reject / Sanitary", "Broken beaker glass, thermocol (EPS), sanitary pads, bandages", "Wrap glass securely in cardboard to protect workers")
    ]

    for row_data in bins_data:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val
            for r in row_cells[i].paragraphs[0].runs:
                r.font.size = Pt(8)

    # Section 4: Test Cases
    doc.add_page_break()
    h4 = doc.add_heading(level=1)
    r4 = h4.add_run("4. Test Cases & Verification Matrix (PPT Benchmark)")
    r4.font.color.rgb = DARK_GREEN

    tc_table = doc.add_table(rows=1, cols=5)
    tc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tc_table.rows[0].cells
    tc_headers = ["Input Item", "Material Detected", "Category & Bin", "Safety Warning", "Status"]
    for i, title_text in enumerate(tc_headers):
        hdr_cells[i].text = title_text
        set_cell_background(hdr_cells[i], "14532D")
        for r in hdr_cells[i].paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8.5)

    tcs = [
        ("Oily samosa wrapper", "Soiled Paper / Food Waste", "Wet / Compost (Green)", "Food oil contamination warning", "PASS"),
        ("Crushed PET soft drink bottle", "PET Plastic #1", "Dry / Recyclable (Blue)", "None (Safe to recycle)", "PASS"),
        ("Electronics lab 9V dead battery", "Alkaline / Zinc Battery", "Hazardous E-Waste (Red)", "Chemical leak & fire hazard", "PASS"),
        ("Broken glass chemistry beaker", "Borosilicate Glass", "Sanitary / Reject (Black)", "Sharp cut hazard for workers", "PASS")
    ]

    for row_data in tcs:
        row_cells = tc_table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val
            for r in row_cells[i].paragraphs[0].runs:
                r.font.size = Pt(8)

    doc.add_paragraph("\nScreenshots from the Live Application:\n")
    img_battery = os.path.join(SCREENSHOTS_DIR, "04_result_battery_hazardous.png")
    if os.path.exists(img_battery):
        doc.add_picture(img_battery, width=Inches(5.5))

    # Section 5: Responsible AI
    doc.add_page_break()
    h5 = doc.add_heading(level=1)
    r5 = h5.add_run("5. Responsible AI Framework")
    r5.font.color.rgb = DARK_GREEN

    doc.add_paragraph(
        "EcoSort Campus enforces six pillars of Responsible AI in its architecture:\n\n"
        "1. Fairness & Inclusion: Calibrated for Indian campus contexts including regional packaging (chai cups, kulhad, samosa paper).\n"
        "2. Transparency & Reasoning: Every response articulates the underlying physical explanation behind bin choices.\n"
        "3. Privacy: Zero personal data retention. No student IDs, photos, device telemetry, or credentials stored.\n"
        "4. Frontline Worker Safety: Proactive hazard guardrails mandate protective handling for broken glass and batteries.\n"
        "5. Institutional Precedence: Local campus environmental protocols always supersede AI guidance.\n"
        "6. Education Over Automation: Teaches sustainable habits by explaining 'why' rather than fostering blind compliance."
    )

    img_resp = os.path.join(SCREENSHOTS_DIR, "09_responsible_ai.png")
    if os.path.exists(img_resp):
        doc.add_picture(img_resp, width=Inches(5.5))

    # Section 6: Projected Impact & Conclusion
    doc.add_page_break()
    h6 = doc.add_heading(level=1)
    r6 = h6.add_run("6. Projected Impact & Future Roadmap")
    r6.font.color.rgb = DARK_GREEN

    doc.add_paragraph(
        "Projected Impact (Concept Pilot Model):\n"
        "\"A pilot projection across a 3,000-student university campus indicates an average 70.2% overall diversion "
        "of municipal solid waste away from local landfills within 60 days of deployment.\"\n"
        "*(Important: This figure represents an estimated model projection from the ideation slide deck, not a measured historical deployment result).*\n\n"
        "Future Enhancements:\n"
        "• Multimodal Computer Vision (IBM Granite Vision image recognition at bin stations)\n"
        "• Multilingual voice and text input (Hindi, Tamil, Telugu)\n"
        "• Ultrasonic IoT fill-level bin sensors\n"
        "• Campus administrative analytics dashboard"
    )

    doc.add_heading("Project Repository & Contact", level=2)
    doc.add_paragraph(
        "• GitHub Repository: https://github.com/pdharv711/Ecosort-campus\n"
        "• Documentation File: docs/EcoSort_Campus_Prototype_Documentation.pdf\n"
        "• Author: Patel Dharv (GitHub: pdharv711)\n"
        "• Program: AI for Sustainability Virtual Internship 2026 (AICTE / 1M1B / IBM)"
    )

    doc.save(DOCX_PATH)
    print(f"[SUCCESS] Generated DOCX documentation: {DOCX_PATH}")

if __name__ == "__main__":
    create_documentation_docx()
